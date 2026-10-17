"""Blind annotation workbooks (identical for annotators A and B) with dropdown validation."""
import csv
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment,Font,PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
ROOT=Path(__file__).resolve().parents[2];OUT=ROOT/'annotation'
TOPICS=['T01_layoffs_unemployment_up','T02_hiring_jobs_growth','T03_prices_rising','T04_inflation_easing','T05_economic_growth','T06_recession_risk','T07_interest_rates_borrowing','T08_housing_costs','T09_gas_prices','T10_food_prices','T11_tariffs_trade','T12_stock_market','T13_other']
ACTORS=['president_admin','congress','fed','democrats','republicans','business','foreign_trade','other','none']
DIRS=['credit','blame','neutral_mention']
CODES={'relevant':['yes','no','unsure'],'tone_overall':['negative','neutral','positive','mixed'],'fact_direction':['bad','good','neutral','mixed','no_fact'],'emotional_intensity':['0','1','2'],
 **{t:['1'] for t in TOPICS},'price_frame':['level','rate','both','none'],'perspective':['national','household','both','none'],'time_frame':['current','forecast','both','none'],
 'numeric_claim':['yes','no'],'claim_text':None,'definition_error':['yes','no','not_applicable'],'actor_1':ACTORS,'direction_1':DIRS,'actor_2':ACTORS,'direction_2':DIRS,'notes':None}
FIXED=['item_id','date','unit_type','context','text']

def sheet(wb,title,rows):
 ws=wb.create_sheet(title);cols=FIXED+list(CODES)
 ws.append(cols)
 for r in rows:ws.append([r[c] for c in FIXED]+['']*len(CODES))
 widths={'item_id':8,'date':11,'unit_type':16,'context':28,'text':80,'claim_text':30,'notes':30}
 for i,c in enumerate(cols,1):
  ws.column_dimensions[get_column_letter(i)].width=widths.get(c,13 if not c.startswith('T') else 6)
  cell=ws.cell(row=1,column=i);cell.font=Font(bold=True);cell.alignment=Alignment(wrap_text=True,vertical='top',text_rotation=90 if c.startswith('T') else 0)
  cell.fill=PatternFill('solid',fgColor='DDDDDD' if c in FIXED else 'FFF2CC')
 ws.row_dimensions[1].height=150
 for row in ws.iter_rows(min_row=2):
  for cell in row:cell.alignment=Alignment(wrap_text=True,vertical='top')
 n=len(rows)+1
 for i,c in enumerate(cols,1):
  vals=CODES.get(c)
  if vals:
   dv=DataValidation(type='list',formula1='"'+','.join(vals)+'"',allow_blank=True,showErrorMessage=True,errorTitle='Invalid code',error='Use a value from the dropdown (see codebook).')
   ws.add_data_validation(dv);dv.add(f'{get_column_letter(i)}2:{get_column_letter(i)}{n}')
 ws.freeze_panes='F2'

def main():
 rows=list(csv.DictReader(open(OUT/'sample_key_DO_NOT_SHOW_ANNOTATORS.csv')))
 calib=[r for r in rows if r['split']=='calibration'];main_rows=[r for r in rows if r['split']!='calibration']
 for who in ['A','B']:
  wb=Workbook();ws=wb.active;ws.title='README'
  for line in [f'Annotator {who}. Read annotation/codebook_v1.md before starting.','1) Label sheet "calibration" (30 items) independently, then discuss with the other annotator and finalize the codebook v1.',
   '2) Label sheet "main" (360 items) independently. Do not discuss items, do not look up sources online.','3) If relevant = no, leave the rest of the row blank.','4) Topic columns T01-T13: put 1 for each topic the text actually discusses.',
   '5) Log start/end times in "time_log".','Do not edit grey columns, sort, or delete rows.']:ws.append([line])
  ws.column_dimensions['A'].width=120
  sheet(wb,'calibration',calib);sheet(wb,'main',main_rows)
  cs=wb.create_sheet('codes');cs.append(['column','allowed values'])
  for c,v in CODES.items():cs.append([c,', '.join(v) if v else 'free text'])
  cs.column_dimensions['A'].width=32;cs.column_dimensions['B'].width=90
  tl=wb.create_sheet('time_log');tl.append(['session','start (local time)','end (local time)','sheet','first item_id','last item_id','notes'])
  for i in range(1,21):tl.append([i])
  wb.save(OUT/f'annotation_workbook_{who}.xlsx')

if __name__=='__main__':main()
