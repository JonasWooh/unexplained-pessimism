"""Inter-annotator agreement for workbooks A/B and an adjudication workbook of disagreements.

Coded dimensions (other than `relevant`) are compared on items both annotators coded
(relevant in {yes, unsure} for both), as the codebook leaves rows blank when relevant = no.
Krippendorff's alpha: two coders, coincidence-matrix form; ordinal metric for
emotional_intensity, nominal elsewhere. Cohen's kappa: unweighted (sklearn).
"""
import csv,json,collections
from pathlib import Path
import openpyxl
from openpyxl.styles import Alignment,Font,PatternFill
from sklearn.metrics import cohen_kappa_score
ROOT=Path(__file__).resolve().parents[2];ANN=ROOT/'annotation';OUT=ANN/'agreement'
import sys;sys.path.insert(0,str(Path(__file__).parent))
from build_workbooks import CODES,TOPICS,ACTORS
NOMINAL=['tone_overall','fact_direction','price_frame','perspective','time_frame','numeric_claim','definition_error','actor_1','direction_1']
ORDINAL=['emotional_intensity']

def load(who,sheet):
 ws=openpyxl.load_workbook(ANN/f'annotation_workbook_{who}.xlsx',read_only=True)[sheet];rows=list(ws.iter_rows(values_only=True));h=rows[0]
 return {r[0]:{c:('' if v is None else str(v).strip()) for c,v in zip(h,r)} for r in rows[1:] if r[0]}

def alpha(pairs,ordinal=False):
 cats=sorted({v for p in pairs for v in p},key=lambda x:(int(x) if ordinal else x))
 if len(cats)<2:return None
 o=collections.Counter()
 for a,b in pairs:o[(a,b)]+=1;o[(b,a)]+=1
 n_c={c:sum(o[(c,k)] for k in cats) for c in cats};n=sum(n_c.values())
 def d2(c,k):
  if not ordinal:return 0.0 if c==k else 1.0
  i,j=sorted((cats.index(c),cats.index(k)))
  return (sum(n_c[cats[g]] for g in range(i,j+1))-(n_c[c]+n_c[k])/2)**2
 Do=sum(o[(c,k)]*d2(c,k) for c in cats for k in cats)
 De=sum(n_c[c]*n_c[k]*d2(c,k) for c in cats for k in cats)/(n-1)
 return None if De==0 else 1-Do/De

def compare(name,pairs,ordinal=False):
 a=[x for x,_ in pairs];b=[y for _,y in pairs]
 agree=sum(x==y for x,y in pairs)/len(pairs) if pairs else None
 kappa=cohen_kappa_score(a,b) if len(set(a)|set(b))>1 else None
 conf=collections.Counter(pairs)
 return dict(dimension=name,n=len(pairs),percent_agreement=agree,cohen_kappa=kappa,krippendorff_alpha=alpha(pairs,ordinal),metric='ordinal' if ordinal else 'nominal',
  categories='|'.join(sorted(set(a)|set(b)))),conf

def actors(r):return {r[k] for k in ['actor_1','actor_2'] if r.get(k) and r[k]!='none'}

def invalid(rows,who,sheet):
 bad=[]
 for iid,r in rows.items():
  for c,vals in CODES.items():
   if vals and r.get(c,'') not in vals+['']:bad.append(dict(annotator=who,sheet=sheet,item_id=iid,column=c,value=r[c]))
  if r['relevant'] in ('yes','unsure') and not r['tone_overall']:bad.append(dict(annotator=who,sheet=sheet,item_id=iid,column='tone_overall',value='blank although relevant'))
  if r['relevant']=='no' and any(r.get(c) for c in CODES if c not in ('relevant','notes')):bad.append(dict(annotator=who,sheet=sheet,item_id=iid,column='*',value='coded although relevant=no'))
 return bad

def run(sheet,key):
 A,B=load('A',sheet),load('B',sheet);ids=sorted(set(A)&set(B))
 if set(A)!=set(B):raise ValueError('item sets differ between workbooks')
 subsets={'all':ids}
 if sheet=='main':
  for split in ['dev','test']:subsets[split]=[i for i in ids if key[i]['split']==split]
  for tier in ['national','local','social']:subsets['tier_'+tier]=[i for i in ids if key[i]['tier']==tier]
 results=[];confs={}
 for sub,sid in subsets.items():
  rel=[(A[i]['relevant'],B[i]['relevant']) for i in sid]
  r,c=compare('relevant',rel);results.append(dict(sheet=sheet,subset=sub,**r));confs[(sheet,sub,'relevant')]=c
  both=[i for i in sid if A[i]['relevant'] in ('yes','unsure') and B[i]['relevant'] in ('yes','unsure')]
  for d in NOMINAL+ORDINAL:
   r,c=compare(d,[(A[i][d] or '(blank)',B[i][d] or '(blank)') for i in both],d in ORDINAL);results.append(dict(sheet=sheet,subset=sub,**r));confs[(sheet,sub,d)]=c
  for t in TOPICS:
   r,c=compare(t,[(A[i][t] or '0',B[i][t] or '0') for i in both]);results.append(dict(sheet=sheet,subset=sub,**r));confs[(sheet,sub,t)]=c
  for act in [x for x in ACTORS if x!='none']:
   r,c=compare('actor_present:'+act,[(str(int(act in actors(A[i]))),str(int(act in actors(B[i])))) for i in both]);results.append(dict(sheet=sheet,subset=sub,**r));confs[(sheet,sub,'actor_present:'+act)]=c
 return A,B,ids,results,confs,invalid(A,'A',sheet)+invalid(B,'B',sheet)

def main():
 OUT.mkdir(exist_ok=True)
 key={r['item_id']:r for r in csv.DictReader(open(ANN/'sample_key_DO_NOT_SHOW_ANNOTATORS.csv'))}
 allres=[];allconf={};bad=[];disagreements=[]
 for sheet in ['calibration','main']:
  A,B,ids,res,confs,inv=run(sheet,key);allres+=res;allconf.update(confs);bad+=inv
  cols=['relevant']+NOMINAL+ORDINAL+TOPICS+['actor_2','direction_2']
  for i in ids:
   diff=[c for c in cols if A[i].get(c,'')!=B[i].get(c,'')]
   if diff:disagreements.append(dict(sheet=sheet,item_id=i,differing=diff,A=A[i],B=B[i]))
 with open(OUT/'agreement_summary.csv','w',newline='') as f:
  w=csv.DictWriter(f,fieldnames=list(allres[0]));w.writeheader();w.writerows(allres)
 with open(OUT/'confusion_matrices.csv','w',newline='') as f:
  w=csv.writer(f);w.writerow(['sheet','subset','dimension','A','B','count'])
  for (sh,sub,d),c in allconf.items():
   for (a,b),n in sorted(c.items()):w.writerow([sh,sub,d,a,b,n])
 with open(OUT/'invalid_entries.csv','w',newline='') as f:
  w=csv.DictWriter(f,fieldnames=['annotator','sheet','item_id','column','value']);w.writeheader();w.writerows(bad)
 # Adjudication workbook: one row per item x differing column, A and B side by side, final left blank.
 # An existing (possibly filled) workbook is never overwritten; it must cover exactly the current disagreements.
 target=OUT/'adjudication_disagreements.xlsx'
 keys={(d['sheet'],d['item_id'],c) for d in disagreements for c in d['differing']}
 if target.exists():
  rows=list(openpyxl.load_workbook(target,read_only=True)['adjudication'].iter_rows(values_only=True))
  have={(r[0],r[1],r[3]) for r in rows[1:]}
  if have!=keys:raise ValueError(f'existing adjudication workbook does not match current disagreements: {len(have^keys)} differing keys')
 else:write_template(disagreements,target)
 summ=dict(items={'calibration':30,'main':360},items_with_any_disagreement={s:sum(1 for d in disagreements if d['sheet']==s) for s in ['calibration','main']},
  adjudication_rows=len(keys),invalid_entries=len(bad))
 (OUT/'agreement_overview.json').write_text(json.dumps(summ,indent=2))
 print(json.dumps(summ))

def write_template(disagreements,target):
 wb=openpyxl.Workbook();ws=wb.active;ws.title='adjudication'
 ws.append(['sheet','item_id','text','column','A','B','final','reason','A_notes','B_notes'])
 for d in disagreements:
  for c in d['differing']:ws.append([d['sheet'],d['item_id'],d['A']['text'],c,d['A'].get(c,''),d['B'].get(c,''),'','',d['A'].get('notes',''),d['B'].get('notes','')])
 for col,wd in zip('ABCDEFGHIJ',[11,8,70,26,14,14,14,30,30,30]):ws.column_dimensions[col].width=wd
 for row in ws.iter_rows():
  for cell in row:cell.alignment=Alignment(wrap_text=True,vertical='top')
 for cell in ws[1]:cell.font=Font(bold=True);cell.fill=PatternFill('solid',fgColor='DDDDDD')
 ws.freeze_panes='D2';wb.save(target)

if __name__=='__main__':main()
