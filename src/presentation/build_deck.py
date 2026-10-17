"""Build the 9-slide presentation and its speaker script. Run from the project root:
    python3 src/presentation/build_deck.py
Reads frozen results through slide_numbers.py and the chart layers from make_slide_figures.py
(run first). Writes presentation/Unexplained_Pessimism_slides.pptx and presentation/SCRIPT.md.
"""
from pathlib import Path
import math
import re
import sys

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.dml import MSO_LINE

sys.path.insert(0, str(Path(__file__).resolve().parent))
import slide_numbers
import script_text
import pptx_fx as fx
from pptx_fx import Beat, Timeline

ROOT = Path(__file__).resolve().parents[2]
FIG = ROOT / 'presentation' / 'build'
OUT = ROOT / 'presentation'

FONT = 'Helvetica Neue'
INK = '1F2328'
MUTED = '6B7280'
FAINT = 'C7CBD1'
CARD = 'F5F6F8'
CHIP = 'ECEEF1'
BLUE = '2A78D6'
ORANGE = 'EB6834'
DARK = '15181C'
WHITE = 'FFFFFF'

W, H = 13.333, 7.5
LM = 0.65                     # left margin, inches


def rgb(h):
    return RGBColor.from_string(h)


def minus(x, d=1):
    return f'{x:.{d}f}'.replace('-', '−')


# ------------------------------------------------------------------ primitives
def text(slide, x, y, w, h, runs, size=16, color=INK, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, name=None, spacing=None, line=1.0):
    """runs: str, or list of paragraphs; a paragraph is a str or a list of (text, overrides) runs."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    paras = runs if isinstance(runs, list) else [runs]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line
        for seg in (para if isinstance(para, list) else [(para, {})]):
            t, o = seg if isinstance(seg, tuple) else (seg, {})
            r = p.add_run()
            r.text = t
            f = r.font
            f.name = FONT
            f.size = Pt(o.get('size', size))
            f.bold = o.get('bold', bold)
            f.italic = o.get('italic', False)
            f.color.rgb = rgb(o.get('color', color))
            sp = o.get('spacing', spacing)
            if sp:
                r._r.get_or_add_rPr().set('spc', str(sp))
        if isinstance(para, list) and para and isinstance(para[0], tuple) and 'space_before' in para[0][1]:
            p.space_before = Pt(para[0][1]['space_before'])
    return tb


def box(slide, x, y, w, h, fill, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08, name=None, line=None):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        s.name = name
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = rgb(fill)
    if line:
        s.line.color.rgb = rgb(line[0])
        s.line.width = Pt(line[1])
        if len(line) > 2:
            s.line.dash_style = line[2]
    else:
        s.line.fill.background()
    st = s._element.find('{http://schemas.openxmlformats.org/presentationml/2006/main}style')
    if st is not None:
        s._element.remove(st)          # no theme effects (shadow) on drawn shapes
    return s


def box_text(s, paras, size=16, color=INK, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
             margins=(0.3, 0.28, 0.3, 0.2)):
    tf = s.text_frame
    tf.word_wrap = True
    tf.margin_left, tf.margin_top, tf.margin_right, tf.margin_bottom = [Inches(m) for m in margins]
    tf.vertical_anchor = anchor
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        for t, o in para:
            r = p.add_run()
            r.text = t
            f = r.font
            f.name = FONT
            f.size = Pt(o.get('size', size))
            f.bold = o.get('bold', bold)
            f.color.rgb = rgb(o.get('color', color))
            if o.get('spacing'):
                r._r.get_or_add_rPr().set('spc', str(o['spacing']))
        if para and 'space_before' in para[0][1]:
            p.space_before = Pt(para[0][1]['space_before'])
    return s


def picture(slide, path, x, y, w, h, name=None):
    p = slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        p.name = name
    return p


def line(slide, x1, y1, x2, y2, color, width=1.5, dash=None, name=None):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(width)
    if dash:
        c.line.dash_style = dash
    if name:
        c.name = name
    return c


def header(slide, kicker, title, num):
    text(slide, LM, 0.42, 10.2, 0.3, kicker.upper(), size=12, color=ORANGE, bold=True, spacing=150,
         name='!!kicker')
    text(slide, LM, 0.8, 12.0, 0.6, title, size=28, color=INK, bold=True, name='!!title', line=0.95)
    text(slide, W - LM - 0.6, H - 0.48, 0.6, 0.25, str(num), size=10, color=MUTED, align=PP_ALIGN.RIGHT,
         name='!!pagenum')


def footer(slide, s):
    return text(slide, LM, H - 0.48, 10.8, 0.3, s, size=9.5, color=MUTED, name='footer')


HYP = [('H1', 'Real pressures', 'High prices and rates the headline numbers miss'),
       ('H2', 'Media', 'Selection, framing or bias in coverage'),
       ('H3', 'Measurement', 'Survey method, timing or the model')]


def chips(slide, active):
    for i, (h, _, _) in enumerate(HYP):
        on = h == active
        s = box(slide, 11.05 + i * 0.6, 0.44, 0.52, 0.34, INK if on else CHIP, radius=0.5, name=f'!!{h}')
        box_text(s, [[(h, {})]], size=11, bold=True, color=WHITE if on else MUTED, align=PP_ALIGN.CENTER,
                 anchor=MSO_ANCHOR.MIDDLE, margins=(0, 0, 0, 0))


CARD_X, CARD_Y, CARD_W, CARD_H, CARD_GAP = LM, 1.85, 3.83, 3.65, 0.3


def cards(slide):
    out = []
    for i, (h, name, desc) in enumerate(HYP):
        s = box(slide, CARD_X + i * (CARD_W + CARD_GAP), CARD_Y, CARD_W, CARD_H, CARD, radius=0.06,
                name=f'!!{h}')
        box_text(s, [[(h, {'size': 15, 'bold': True, 'color': ORANGE, 'spacing': 100})],
                     [(name, {'size': 26, 'bold': True, 'space_before': 4})],
                     [(desc, {'size': 15, 'color': MUTED, 'space_before': 6})]],
                 margins=(0.32, 0.3, 0.32, 0.2))
        out.append(s)
    return out


# ------------------------------------------------------------------ slides
def s1_title(prs, N):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = rgb(DARK)
    picture(s, FIG / 'title_lines.png', 0, 4.9, W, 2.6)
    text(s, LM + 0.15, 1.45, 11, 0.35, 'AI APPLICATIONS IN FINANCE  ·  FINAL PROJECT', size=13, color=ORANGE,
         bold=True, spacing=200)
    text(s, LM + 0.15, 1.9, 11.5, 1.1, 'Unexplained Pessimism', size=60, color=WHITE, bold=True)
    text(s, LM + 0.15, 3.1, 10.5, 0.9,
         'Economic fundamentals, media narratives and U.S. consumer sentiment before the 2026 midterms',
         size=22, color='C9CDD3', line=1.1)
    text(s, LM + 0.15, 4.2, 10, 0.4, 'Jonas Wu   ·   Data through October 16, 2026', size=15, color='9AA1AB')
    return s, Timeline()


def s2_gap(prs, N):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    header(s, 'The puzzle', f'Sentiment is ~{round(-N["gap_2026"])} points below the economy’s prediction', 2)
    x, y, w, h = 0.35, 1.7, 8.7, 4.95
    picture(s, FIG / 'gap_base.png', x, y, w, h)
    model = picture(s, FIG / 'gap_model.png', x, y, w, h, name='layer model')
    actual = picture(s, FIG / 'gap_actual.png', x, y, w, h, name='layer actual')
    gap = picture(s, FIG / 'gap_gap.png', x, y, w, h, name='layer gap')

    def stat(yy, big, small, color):
        return text(s, 9.55, yy, 3.2, 1.3, [[(big, {'size': 46, 'bold': True, 'color': color})],
                                             [(small, {'size': 14, 'color': MUTED, 'space_before': 2})]])
    st1 = stat(1.95, minus(N['gap_2025']), 'points, 2025 average gap', ORANGE)
    st2 = stat(3.45, minus(N['gap_2026']), 'points, 2026 average gap', ORANGE)
    st3 = stat(4.95, f'{round(100 * N["explained_lo"])}–{round(100 * N["explained_hi"])}%',
               'of the fall since 2019 that the economy explains', INK)
    footer(s, 'Model: real-time unemployment, PCE inflation and GDP growth (ALFRED vintages), fit 2010–19. '
              'Band: 95% prediction interval. Pandemic months not modelled.')
    tl = Timeline().click(Beat(model, 'fade', 600)).click(Beat(actual, 'wipe_left', 1200)) \
        .click(Beat(gap, 'fade', 700), Beat(st1, 'rise', 450), Beat(st2, 'rise', 450),
               Beat(st3, 'rise', 450, delay=150))
    return s, tl


def s3_three(prs, N):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    header(s, 'Three explanations', 'Three ways to explain the gap', 3)
    cs = cards(s)
    qs = []
    for i in range(3):
        cx = CARD_X + i * (CARD_W + CARD_GAP)
        qs.append(text(s, cx, CARD_Y + 2.2, CARD_W, 1.0, '?', size=60, color=FAINT, bold=True,
                       align=PP_ALIGN.CENTER, name=f'verdict q{i}'))
    strip = text(s, LM, 6.0, W - 2 * LM, 0.4,
                 [[('●  ', {'color': ORANGE, 'size': 12}), ('All tests fixed before estimation', {}),
                   ('        ●  ', {'color': ORANGE, 'size': 12}), ('150 specifications registered', {}),
                   ('        ●  ', {'color': ORANGE, 'size': 12}), ('8 pre-set tests against the media story', {})]],
                 size=17, align=PP_ALIGN.CENTER)
    tl = Timeline().auto(*[Beat([c, q], 'rise', 500, delay=0 if i == 0 else 100)
                           for i, (c, q) in enumerate(zip(cs, qs))]) \
        .click(Beat(strip, 'fade', 600))
    return s, tl


def tile(s, x, y, w, big, small, color=INK):
    t = text(s, x, y, w, 1.35, [[(big, {'size': 34, 'bold': True, 'color': color})],
                                [(small, {'size': 14, 'color': MUTED, 'space_before': 3})]], line=1.0)
    return t


def s4_measure(prs, N):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    header(s, 'H3  ·  Measurement', 'Not the survey, not party mix, not the campaign', 4)
    chips(s, 'H3')
    picture(s, FIG / 'groups_base.png', 0.3, 1.6, 6.9, 5.0)
    x = 7.75
    bar = lambda yy: box(s, x, yy + 0.08, 0.06, 1.05, ORANGE, shape=MSO_SHAPE.RECTANGLE)
    t1 = tile(s, x + 0.3, 1.85, 4.7, f'{minus(N["web_gap_start"], 0)} → {minus(N["web_gap_end"], 0)}',
              'Gap kept widening after the survey moved online (late 2024 → 2026)')
    t2 = tile(s, x + 0.3, 3.4, 4.7, f'≤ {N["party_bound"]:.1f} pts',
              'Largest plausible effect of the party mix of respondents')
    t3 = tile(s, x + 0.3, 4.95, 4.7, f'{minus(N["gap_2026_janapr"])} → {minus(N["gap_2026_mayoct"])}',
              'Gap in Jan–Apr vs. May–Oct 2026: already there before the campaign')
    b1, b2, b3 = bar(1.85), bar(3.4), bar(4.95)
    footer(s, 'Group gaps: each group’s own economic model, fit 2010–19; mean 2025-01 to 2026-08. '
              'Party-mix bound: Republican share 2–6 points lower.')
    tl = Timeline().click(Beat([b1, t1], 'rise', 500)).click(Beat([b2, t2], 'rise', 500)) \
        .click(Beat([b3, t3], 'rise', 500))
    return s, tl


def s5_ai(prs, N):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    header(s, 'H2  ·  Measuring the media', 'An AI coder that nearly matches human agreement', 5)
    chips(s, 'H2')
    nodes = [(f'{int(N["corpus_items"]):,}', 'news items, posts and comments'),
             (f'{int(N["coded_items"]):,}', 'coded by an LLM'),
             (f'{int(N["gold_items"])}', 'labeled by two people, independently')]
    nw, gapw = 3.45, 0.6
    seq = []
    for i, (big, small) in enumerate(nodes):
        nx = LM + i * (nw + gapw)
        b = box(s, nx, 1.75, nw, 1.2, CARD, radius=0.1)
        box_text(b, [[(big, {'size': 30, 'bold': True})], [(small, {'size': 13.5, 'color': MUTED})]],
                 margins=(0.28, 0.17, 0.2, 0.1))
        seq.append(b)
        if i < 2:
            a = text(s, nx + nw, 1.75, gapw, 1.2, '→', size=26, color=FAINT, align=PP_ALIGN.CENTER,
                     anchor=MSO_ANCHOR.MIDDLE)
            seq.append(a)
    cap = text(s, LM, 3.3, 8.5, 0.35, 'Agreement with the human labels on tone (Cohen’s κ, test set)', size=13,
               color=MUTED, bold=True)
    rows = [('Human vs. human', N['kappa_human'], INK), ('LLM', N['kappa_llm'], ORANGE),
            ('TF-IDF classifier', N['kappa_tfidf'], FAINT), ('Finance dictionary', N['kappa_lm'], FAINT),
            ('VADER', N['kappa_vader'], FAINT)]
    bx, scale, bars = 3.05, 6.6, []
    for i, (lab, k, col) in enumerate(rows):
        yy = 3.8 + i * 0.58
        l1 = text(s, LM, yy, bx - LM - 0.2, 0.4, lab, size=15, align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE,
                  bold=(lab == 'LLM'))
        r = box(s, bx, yy + 0.03, k * scale, 0.34, col, shape=MSO_SHAPE.RECTANGLE)
        v = text(s, bx + k * scale + 0.12, yy, 0.8, 0.4, f'{k:.2f}', size=15, bold=True,
                 color=INK if col == FAINT else col, anchor=MSO_ANCHOR.MIDDLE)
        bars.append((l1, r, v))
    badge = box(s, 10.05, 4.05, 2.65, 2.0, CARD, radius=0.1)
    box_text(badge, [[('✓', {'size': 30, 'bold': True, 'color': ORANGE})],
                     [('Name and source swaps', {'size': 14, 'bold': True})],
                     [(f'{int(N["swap_sig"])} of {int(N["swap_tests"])} tests show bias',
                       {'size': 13, 'color': MUTED, 'space_before': 2})]],
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, margins=(0.15, 0.1, 0.15, 0.1))
    footer(s, f'Gold standard: {int(N["gold_items"])} items, two annotators, adjudicated; LLM prompt frozen on '
              f'100 development items and scored once on 260 test items.')
    tl = Timeline().auto(*[Beat(x, 'fade', 400, delay=0 if i == 0 else 50) for i, x in enumerate(seq)]) \
        .click(Beat(cap, 'fade', 300), *[Beat([l, r, v], 'wipe_left', 550) for l, r, v in bars]) \
        .click(Beat(badge, 'rise', 500))
    return s, tl


def s6_coverage(prs, N):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    header(s, 'H2  ·  What coverage shows', 'Even-handed tone, but more blame on the White House', 6)
    chips(s, 'H2')
    checks = ['Good and bad news treated alike', 'Rising and easing prices treated alike',
              f'{int(N["fact_wrong"])} of {int(N["fact_checkable"])} checkable numbers wrong']
    items = []
    for i, c in enumerate(checks):
        items.append(text(s, LM, 2.2 + i * 1.05, 5.4, 0.6,
                          [[('✓   ', {'color': ORANGE, 'bold': True, 'size': 22}), (c, {})]],
                          size=19, anchor=MSO_ANCHOR.MIDDLE))
    cap = text(s, 6.9, 1.85, 5.8, 0.35, 'Share of economic news items blaming the White House', size=13,
               color=MUTED, bold=True)
    admins = [('Obama', N['blame_obama']), ('Trump I', N['blame_trump1']), ('Biden', N['blame_biden']),
              ('Trump II', N['blame_trump2'])]
    base_y, top_h, bw, step = 5.75, 3.2, 0.95, 1.4
    bars = []
    for i, (a, v) in enumerate(admins):
        bx = 7.15 + i * step
        hh = v / 0.45 * top_h
        col = ORANGE if a == 'Trump II' else FAINT
        r = box(s, bx, base_y - hh, bw, hh, col, shape=MSO_SHAPE.RECTANGLE)
        vl = text(s, bx - 0.3, base_y - hh - 0.48, bw + 0.6, 0.4, f'{round(100 * v)}%', size=20, bold=True,
                  color=ORANGE if a == 'Trump II' else INK, align=PP_ALIGN.CENTER)
        lb = text(s, bx - 0.3, base_y + 0.1, bw + 0.6, 0.35, a, size=15, color=INK, align=PP_ALIGN.CENTER)
        bars.append((r, vl, lb))
    axis = line(s, 6.95, base_y, 12.7, base_y, FAINT, 1.25)
    note = text(s, 6.9, 6.35, 5.9, 0.45,
                [[('Blame is not proof of bias: ', {'bold': True}), ('tariffs are the president’s own decisions', {})]],
                size=14)
    footer(s, f'National news (NYT, Guardian), LLM-coded. Factual check against the official data available on '
              f'the publication date.')
    tl = Timeline().auto(*[Beat(x, 'rise', 450, delay=0 if i == 0 else 100) for i, x in enumerate(items)]) \
        .click(Beat([cap, axis], 'fade', 300),
               *[Beat([r, lb], 'wipe_up', 500) for r, _, lb in bars[:3]],
               Beat([v for _, v, _ in bars[:3]], 'fade', 300),
               Beat([bars[3][0], bars[3][2]], 'wipe_up', 800), Beat(bars[3][1], 'fade', 400)) \
        .click(Beat(note, 'rise', 500))
    return s, tl


def s7_forecast(prs, N):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    header(s, 'H2  ·  The forecast test', 'Tone stops helping forecasts once controls are in', 7)
    chips(s, 'H2')
    x, y, w, h = 0.75, 1.75, 11.6, 3.55
    picture(s, FIG / 'spec_base.png', x, y, w, h)
    r1 = picture(s, FIG / 'spec_row1.png', x, y, w, h, name='layer row1')
    r2 = picture(s, FIG / 'spec_row2.png', x, y, w, h, name='layer row2')
    big = text(s, LM + 0.2, 5.55, 2.6, 0.95, f'≤ {math.ceil(100 * N["tone_share_max"])}%', size=48, bold=True,
               color=ORANGE, anchor=MSO_ANCHOR.MIDDLE)
    lab = text(s, LM + 2.75, 5.55, 8.5, 0.95, 'of the 2025–26 gap is linked to news tone, at most', size=20,
               color=INK, anchor=MSO_ANCHOR.MIDDLE)
    footer(s, 'One-step-ahead forecasts 2017–26, Clark–West test. Outlined dot = main specification. '
              'Controls: household pressures, governing party, pre-election month.')
    tl = Timeline().click(Beat(r1, 'fade', 600)).click(Beat(r2, 'fade', 600)) \
        .click(Beat([big, lab], 'rise', 500))
    return s, tl


def s8_verdict(prs, N):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    header(s, 'Verdict', 'Media bias is not the main explanation', 8)
    cards(s)
    verdicts = {'H1': ('OPEN', BLUE, 'Plausible, but the 2010–19 data cannot size it'),
                'H2': ('WEAK', ORANGE, f'≤ {math.ceil(100 * N["tone_share_max"])}% of the gap; 2 of 8 tests go against it'),
                'H3': ('RULED OUT', INK, 'Not survey mode, party mix or timing')}
    beats = {}
    for i, (h, _, _) in enumerate(HYP):
        cx = CARD_X + i * (CARD_W + CARD_GAP)
        lab, col, why = verdicts[h]
        pill = box(s, cx + 0.32, CARD_Y + 2.1, 1.9 if lab != 'RULED OUT' else 2.2, 0.5, col, radius=0.5)
        box_text(pill, [[(lab, {'size': 14, 'bold': True, 'color': WHITE, 'spacing': 100})]],
                 align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, margins=(0, 0, 0, 0))
        reason = text(s, cx + 0.32, CARD_Y + 2.75, CARD_W - 0.6, 0.8, why, size=14, color=INK, line=1.05)
        beats[h] = (pill, reason)
    foot = text(s, LM, 5.95, W - 2 * LM, 0.7,
                [[('Limits  ', {'bold': True, 'color': INK}),
                  ('observational data; two national outlets, about 26 items a month', {})],
                 [('Next  ', {'bold': True, 'color': INK}),
                  ('a randomized framing experiment to test cause', {'space_before': 4})]],
                size=14, color=MUTED, align=PP_ALIGN.LEFT)
    tl = Timeline()
    for h in ('H3', 'H2', 'H1'):
        p, r = beats[h]
        if h == 'H1':
            tl.click(Beat(p, 'rise', 450), Beat(r, 'fade', 350), Beat(foot, 'fade', 500, delay=400))
        else:
            tl.click(Beat(p, 'rise', 450), Beat(r, 'fade', 350))
    return s, tl


def s9_implication(prs, N):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    header(s, 'Implication for forecasting and risk', 'Treat the gap as persistent, not as noise', 9)
    cap = text(s, LM, 2.05, 6.4, 0.35, 'Michigan index, 2026 average', size=13, color=MUTED, bold=True)
    x0, scale = LM, 0.066
    rows = [('Economy predicts', N['pred_2026'], BLUE, 2.65), ('People say', N['actual_2026'], INK, 4.0)]
    shapes = []
    for lab, v, col, yy in rows:
        l = text(s, x0, yy, 4, 0.35, lab, size=15, bold=True, color=col)
        b = box(s, x0, yy + 0.42, v * scale, 0.55, col, shape=MSO_SHAPE.RECTANGLE)
        n = text(s, x0 + v * scale + 0.15, yy + 0.38, 1.2, 0.6, f'{round(v)}', size=28, bold=True, color=col,
                 anchor=MSO_ANCHOR.MIDDLE)
        shapes.append((l, b, n))
    mx = x0 + N['pre2020_min'] * scale
    mk = line(s, mx, 2.5, mx, 5.25, ORANGE, 1.75, dash=MSO_LINE.DASH)
    mkl = text(s, mx - 2.6, 5.35, 5.2, 0.6,
               [[(f'Lowest reading before 2020: {N["pre2020_min"]:.1f}', {'bold': True, 'color': ORANGE})],
                [(f'({N["pre2020_min_when"]}; series since {int(N["michigan_first_year"])})',
                  {'color': MUTED, 'size': 12})]], size=13, align=PP_ALIGN.CENTER)
    takeaways = [('Don’t use pre-2020 mappings', 'They would read today as a deep downturn'),
                 ('Track the gap and its error band', 'Next to price levels and mortgage rates'),
                 ('Skip news-tone signals here', 'No forecast gain once those are included')]
    tk = []
    for i, (a, b) in enumerate(takeaways):
        yy = 2.2 + i * 1.45
        num = text(s, 7.7, yy, 0.6, 0.8, str(i + 1), size=34, bold=True, color=ORANGE)
        body = text(s, 8.35, yy + 0.05, 4.4, 1.0, [[(a, {'size': 18, 'bold': True})],
                                                   [(b, {'size': 14, 'color': MUTED, 'space_before': 3})]])
        tk.append((num, body))
    footer(s, 'Prediction: economy-only model (real-time data, fit 2010–19), January–October 2026. '
              'The dashboard tracks the gap and its band each month.')
    tl = Timeline().click(Beat(cap, 'fade', 300), Beat([shapes[0][0], shapes[0][1]], 'wipe_left', 900),
                          Beat(shapes[0][2], 'fade', 300),
                          Beat([shapes[1][0], shapes[1][1]], 'wipe_left', 700), Beat(shapes[1][2], 'fade', 300)) \
        .click(Beat([mk, mkl], 'fade', 500)) \
        .click(*[Beat(list(t), 'rise', 450, delay=0 if i == 0 else 150) for i, t in enumerate(tk)])
    return s, tl


SLIDES = [s1_title, s2_gap, s3_three, s4_measure, s5_ai, s6_coverage, s7_forecast, s8_verdict, s9_implication]
TRANSITION = [None, 'fade', 'morph', 'morph', 'morph', 'morph', 'morph', 'morph', 'morph']


def words(s):
    return len(re.sub(r'\[CLICK\]', ' ', s).split())


def main():
    N, _ = slide_numbers.compute()
    # wording in the script that summarizes a number must stay true to it
    assert 0.05 <= N['explained_lo'] and N['explained_hi'] <= 0.15, 'script says "about a tenth"'
    assert 1.8 <= N['kappa_llm'] / N['kappa_lm'] <= 2.3, 'script says "about twice as well"'
    assert N['actual_2026'] < N['pre2020_min'], 'script says "lower than any reading before 2020"'
    assert N['m3_n_sig'] == 0 and N['party_bound'] < 3
    sc = script_text.script(N)

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(W), Inches(H)
    rows = []
    for i, (fn, tr, sp) in enumerate(zip(SLIDES, TRANSITION, sc)):
        slide, tl = fn(prs, N)
        n_click = sp['en'].count('[CLICK]')
        assert n_click == tl.n_clicks, f'slide {i + 1}: script has {n_click} clicks, animation {tl.n_clicks}'
        fx.apply(slide, transition=tr, dur_ms=900 if tr == 'morph' else 700, timeline=tl)
        nw = words(sp['en'])
        secs = nw / 130 * 60 + n_click * 0.8
        rows.append((i + 1, nw, n_click, secs))
        slide.notes_slide.notes_text_frame.text = (
            f'≈ {round(secs)} s  ·  {nw} words  ·  {n_click} clicks\n\n{sp["en"]}')
    out = OUT / 'Unexplained_Pessimism_slides.pptx'
    prs.save(out)

    tot_w = sum(r[1] for r in rows)
    tot_c = sum(r[2] for r in rows)
    t = lambda wpm: tot_w / wpm * 60 + tot_c * 0.8 + 8 * 1.0
    ms = lambda sec: f'{int(sec // 60)}:{round(sec % 60):02d}'
    md = ['# Presentation script', '',
          f'Slides: `Unexplained_Pessimism_slides.pptx` (9 slides). [CLICK] = press the advance key once to trigger '
          f'the next animation step; after a slide is finished, press once more to go to the next slide.', '',
          f'**Length**: {tot_w} words, {tot_c} clicks. Estimated total (including about 0.8 s per click '
          f'and 1 s per slide change): {ms(t(140))} at 140 words per minute, {ms(t(130))} at 130, '
          f'{ms(t(120))} at 120. Limit 4:30.', '']
    md += ['| Slide | Words | Clicks | ≈ seconds at 130 wpm |', '| --- | --- | --- | --- |']
    md += [f'| {a} | {b} | {c} | {d:.0f} |' for a, b, c, d in rows]
    md.append('')
    titles = ['Title', 'The puzzle', 'Three explanations', 'H3 · Measurement', 'H2 · Measuring the media',
              'H2 · What coverage shows', 'H2 · The forecast test', 'Verdict', 'Implication']
    for (i, nw, nc, secs), sp, ttl in zip(rows, sc, titles):
        md += [f'## {i}. {ttl}  (≈ {secs:.0f} s)', '', sp['en'], '']
    (OUT / 'SCRIPT.md').write_text('\n'.join(md))
    print(out)
    for r in rows:
        print(r)
    print('total words', tot_w, 'clicks', tot_c, 'est s @140/130/120:', round(t(140)), round(t(130)), round(t(120)))


if __name__ == '__main__':
    main()
